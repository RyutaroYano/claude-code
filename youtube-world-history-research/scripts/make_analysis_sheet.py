"""はねた動画の分析シートを xlsx と csv で出力する。

引数なし: 4項目が空欄の記入用シート
--filled: analysis_filled.py の考察で4項目を埋めたシート

動画タイトルは vidIQ で取得した値（2026-10-08〜09時点）。
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CHANNELS = {
    "世界史裏探訪": "https://www.youtube.com/@sekaishi_ura",
    "おもしろ世界史ラボ": "https://www.youtube.com/@%E3%81%8A%E3%82%82%E3%81%97%E3%82%8D%E4%B8%96%E7%95%8C%E5%8F%B2%E3%83%A9%E3%83%9C",
    "追伸、戦場から失礼します。": "https://www.youtube.com/@hello-senjo",
}

# (チャンネル名, 動画ID, 動画タイトル)
VIDEOS = [
    ("世界史裏探訪", "y0cwR7ttIaY", "なぜ最強だったはずのモンゴル帝国は ヨーロッパを侵略できなかったのか？【歴史解説】"),
    ("世界史裏探訪", "t_Kje4VAKiE", "【立場逆転】侵略しに来たのに狩られる側に…モンゴル目線で見る元寇の恐怖【歴史解説】"),
    ("世界史裏探訪", "62_F5SjtZCM", "【恐怖】日本のバケモン海賊「倭寇」、さすがに厄介すぎる【歴史解説】"),
    ("世界史裏探訪", "e-i5fvSbPM0", "街角の子供から戦場の兵士まで...世界を支配したコカ・コーラの歴史【歴史解説】"),
    ("世界史裏探訪", "_CE6EZ7nBWY", "【総集編】犯罪者が建国し、凡ミスで滅亡...歴史上存在したヤバすぎる国たち【歴史解説】"),
    ("世界史裏探訪", "lOGdA7_YDe0", "【致〇率99%】敵陣に最初に突入する兵士のモチベーションは何だったのか？【歴史解説】"),
    ("世界史裏探訪", "EXsFtJWLtLw", "ビックリするほどしょーもない理由で滅亡した国３選【歴史解説】"),
    ("世界史裏探訪", "oqU0VDyQiUE", "なぜ攻撃力最強のはずの機関銃は、軍人たちから嫌われまくっていたのか？【歴史解説】"),
    ("世界史裏探訪", "mMzDpd9NA38", "「歩兵」が中世ヨーロッパの戦場を”地獄”に変えたと言われる理由【歴史解説】"),
    ("世界史裏探訪", "ebAl58P2JRc", "防御力最強のはずの「城塞都市」が日本で採用されなかった意外な理由【歴史解説】"),
    ("おもしろ世界史ラボ", "sky5rp3O0fg", "なぜインドは「人口世界一」になったのか？14億人を生んだ最強の地形と残酷な歴史"),
    ("おもしろ世界史ラボ", "_rEy7JfXiKE", "【驚愕の地理】日本より細長い国が存在した！？チリ・ノルウェー・ベトナムが抱える過酷な宿命と驚きの共通点"),
    ("おもしろ世界史ラボ", "b8Hy-06K0mY", "【大人の歴史物語】なぜ北米の巨大な島「グリーンランド」は、遠く離れたデンマーク領なのか？ 〜嘘と勘違いから始まった数奇な運命〜"),
    ("おもしろ世界史ラボ", "0fQNhKQyVcg", "【地政学と歴史】中国はなぜ日本海に出られないのか？ロシアに奪われた「15kmの壁」の真実"),
    ("おもしろ世界史ラボ", "L1HApBYXYaw", "なぜ中央アジアは「5つのスタン」に分かれたのか？【シルクロードの真実】"),
    ("おもしろ世界史ラボ", "YgU5vvbNHFc", "ロシア最大の痛恨…なぜ「世界最大の宝の山」を手放したのか？坪単価0.2円で売られたアラスカの悲劇と歴史の皮肉"),
    ("おもしろ世界史ラボ", "WTxwe4d4qXM", "一見すると双子なのに…オランダとベルギーはなぜ「別の国」になったのか？【引き裂かれた歴史】"),
    ("おもしろ世界史ラボ", "rJqy3SXia3Y", "国旗も言葉もほぼ同じなのに…北欧が「一つの国」になれなかった理由"),
    ("おもしろ世界史ラボ", "dkrkDvQjz-4", "一つ屋根の下に二つの家族？イベリア半島が「２つに分かれた」歴史の裏側"),
    ("おもしろ世界史ラボ", "MIE3-uo--r4", "同じ言語なのに、なぜ別の国？ドイツとオーストリアの複雑すぎる愛憎史"),
    ("追伸、戦場から失礼します。", "BtAGQokhyUE", "ヘルマン・ゲーリング｜残酷すぎる最期。53歳で死去したヒトラーの後継者【第二次世界大戦】"),
    ("追伸、戦場から失礼します。", "XBXVSbh4w1k", "イルマ・グレーゼ｜22歳で処刑された悪魔。ナチスが最も恐れた女の残酷すぎる最期【第二次世界大戦】"),
    ("追伸、戦場から失礼します。", "K4AcedCAqj4", "ラインハルト・ハイドリヒ｜数百万人を虐殺した悪魔。ナチスが最も恐れた男の残酷すぎる最期【第二次世界大戦】"),
    ("追伸、戦場から失礼します。", "UdZXaobRj_U", "毛沢東の最期！4500万人を殺した男の残酷すぎる44日間【第二次世界大戦】"),
    ("追伸、戦場から失礼します。", "6KhwEtMg3-M", "ラヴレンチー・ベリヤ｜数百万人を虐殺した悪魔。スターリンが絶賛したソ連の秘密警察長官【第二次世界大戦】"),
    ("追伸、戦場から失礼します。", "b0K8lXRPE_g", "エルヴィン・ロンメル｜ヒトラーに殺された男。伝説の無敵将校「砂漠の狐」の残酷すぎる最期【第二次世界大戦】"),
    ("追伸、戦場から失礼します。", "bHrzNTlZntI", "アドルフ・アイヒマン｜500万人を殺した男。ナチスが最も恐れた男の残酷すぎる最期【第二次世界大戦】"),
    ("追伸、戦場から失礼します。", "Rqgbenusa_8", "スターリンの最期！1000万人を殺した男の残酷すぎる7日間【第二次世界大戦】"),
    ("追伸、戦場から失礼します。", "ANxbHSdWPUg", "ムッソリーニの最期！逆さ吊りで公開処刑されるまでの残酷すぎる10日間【第二次世界大戦】"),
    ("追伸、戦場から失礼します。", "PINVVV6JeGg", "ヒトラーの最期！最恐独裁者の悲惨すぎる最後の48時間【第二次世界大戦】"),
]

HEADERS = [
    "チャンネル名", "チャンネルURL", "動画タイトル", "動画URL",
    "誰に向けた動画か", "見たら何が得られると約束しているか",
    "どんな感情を動かしているか（驚き・不安・好奇心など）", "サムネの構図・文字数・色・人物の表情",
]
WIDTHS = [22, 40, 60, 44, 30, 30, 30, 36]


def rows(analysis=None):
    for ch, vid, title in VIDEOS:
        filled = list(analysis[vid]) if analysis else ["", "", "", ""]
        yield [ch, CHANNELS[ch], title, f"https://www.youtube.com/watch?v={vid}", *filled]


def main():
    filled = "--filled" in sys.argv
    analysis = None
    name = "video_analysis_sheet"
    if filled:
        from analysis_filled import ANALYSIS as analysis
        name = "video_analysis_sheet_filled"
    out_dir = os.path.join(os.path.dirname(__file__), "..", "output")

    wb = Workbook()
    ws = wb.active
    ws.title = "動画分析"
    ws.append(HEADERS)
    for r in rows(analysis):
        ws.append(r)

    header_fill = PatternFill("solid", fgColor="D9E2F3")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row[4:]:
            cell.fill = input_fill
        for cell in (row[1], row[3]):
            cell.hyperlink = cell.value
            cell.font = Font(color="0563C1", underline="single")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    wb.save(os.path.join(out_dir, f"{name}.xlsx"))

    with open(os.path.join(out_dir, f"{name}.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(rows(analysis))


if __name__ == "__main__":
    main()
